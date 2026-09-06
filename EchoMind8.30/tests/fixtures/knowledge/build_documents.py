"""生成多格式知识库解析测试夹具。"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


ROOT = Path(__file__).parent
FONT_PATH = "C:/Windows/Fonts/msyh.ttc"
FONT_NAME = "MicrosoftYaHei"


def set_run_font(run, size=11, bold=False):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold


def add_paragraph(document, text, style=None):
    paragraph = document.add_paragraph(style=style)
    set_run_font(paragraph.add_run(text))
    return paragraph


def create_docx():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run("商城退款与售后规则"), size=18, bold=True)
    intro = add_paragraph(document, "本规则用于说明订单退款、退货与人工升级的处理条件，供客服人员统一答复用户咨询。")
    intro.paragraph_format.space_after = Pt(10)

    add_paragraph(document, "一 退款申请条件", style="Heading 1")
    add_paragraph(document, "用户可在订单签收后七日内提交无理由退款申请。商品应保持完整，配件、赠品和包装应一并退回。")
    add_paragraph(document, "一 无理由退款", style="Heading 2")
    add_paragraph(document, "退款申请提交后，系统通常在一至三个工作日内完成审核；审核通过后，款项按原支付方式退回。")
    add_paragraph(document, "二 质量问题退款", style="Heading 2")
    add_paragraph(document, "如商品存在质量问题，用户应提供订单号、问题照片或视频。确认后可按平台规则安排退货或退款。")

    add_paragraph(document, "二 审核与到账时效", style="Heading 1")
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["处理阶段", "预计时效", "客服说明"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = ""
        set_run_font(cell.paragraphs[0].add_run(text), bold=True)
    rows = [
        ("退款审核", "1 至 3 个工作日", "审核期间可在订单详情查看进度"),
        ("原路退款", "5 至 7 个工作日", "到账时间受支付机构处理周期影响"),
        ("退货物流核验", "签收后 1 个工作日", "请保留寄回物流单号"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = ""
            set_run_font(cell.paragraphs[0].add_run(value))

    add_paragraph(document, "三 人工升级条件", style="Heading 1")
    add_paragraph(document, "出现重复扣款、退款超时、商品质量争议或用户明确要求人工处理时，客服应创建人工升级工单，并记录订单号与问题描述。")
    document.save(ROOT / "商城退款与售后规则.docx")


def create_text_pdf():
    pdfmetrics.registerFont(TTFont(FONT_NAME, FONT_PATH, subfontIndex=0))
    path = ROOT / "技术支持与错误码手册.pdf"
    pdf = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    y = height - 25 * mm

    def line(text, size=10, leading=7, color=None):
        nonlocal y
        pdf.setFont(FONT_NAME, size)
        pdf.setFillColor(color or HexColor("#202020"))
        pdf.drawString(22 * mm, y, text)
        y -= leading * mm

    line("技术支持与错误码处理手册", size=18, leading=12)
    line("本手册用于处理登录、页面加载、支付与服务异常等技术咨询。", size=10, leading=10)
    line("一 受理信息", size=14, leading=9, color=HexColor("#1F4E79"))
    line("请优先收集：用户标识、应用版本、设备系统、网络环境、错误码、发生时间和问题截图。")
    line("二 常见错误码", size=14, leading=9, color=HexColor("#1F4E79"))
    line("错误码 401 认证失效", size=11, leading=7)
    line("含义：登录凭证失效或账户认证失败。处理：引导用户重新登录；持续失败时建议重置密码。")
    line("错误码 403 无访问权限", size=11, leading=7)
    line("含义：当前账户无权访问对应资源。处理：确认账户状态与权限范围，不要要求用户提供密码。")
    line("错误码 500 服务端异常", size=11, leading=7)
    line("含义：服务端处理失败。处理：建议稍后重试；如持续发生，记录错误码与时间后升级技术工单。")
    line("三 排查步骤", size=14, leading=9, color=HexColor("#1F4E79"))
    line("1 检查网络：切换 WiFi 与移动网络后重新操作。")
    line("2 检查版本：升级到最新应用版本，清除缓存后重启。")
    line("3 检查服务：如存在服务故障公告，向用户说明影响范围与恢复进度。")
    line("四 人工升级", size=14, leading=9, color=HexColor("#1F4E79"))
    line("发生数据丢失、支付成功但订单未生成、重复出现 500 错误或用户明确要求人工处理时，创建技术支持工单。")
    pdf.save()


def create_scanned_pdf():
    path = ROOT / "物流异常处理指南_扫描版.pdf"
    pages = [
        [
            "物流异常处理指南",
            "一 物流信息停滞",
            "发货后 24 小时内未更新属于正常范围。",
            "超过 48 小时未更新时，收集订单号和物流单号。",
            "二 包裹超过预计时效未送达",
            "先确认收货地址和物流轨迹，再提交查件申请。",
        ],
        [
            "物流异常处理指南",
            "三 显示已签收但未收到",
            "建议用户联系配送员并核对代收情况。",
            "仍未找到包裹时，记录异常并升级人工客服。",
            "四 地址修改",
            "订单未发货时可申请修改；发货后需联系承运方确认。",
        ],
    ]
    images = []
    for index, lines in enumerate(pages, start=1):
        image = Image.new("RGB", (1654, 2339), "#F6F1E5")
        draw = ImageDraw.Draw(image)
        title_font = ImageFont.truetype(FONT_PATH, 58, index=0)
        body_font = ImageFont.truetype(FONT_PATH, 34, index=0)
        y = 180
        for line_index, text in enumerate(lines):
            font = title_font if line_index == 0 else body_font
            color = "#1F4E79" if line_index in {0, 1, 4} else "#222222"
            draw.text((150, y), text, font=font, fill=color)
            y += 130 if line_index == 0 else 95
        image_path = ROOT / f".scan_page_{index}.png"
        image.save(image_path)
        images.append(image_path)

    pdf = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    for image_path in images:
        pdf.drawImage(str(image_path), 0, 0, width=width, height=height)
        pdf.showPage()
        image_path.unlink()
    pdf.save()


if __name__ == "__main__":
    ROOT.mkdir(parents=True, exist_ok=True)
    create_docx()
    create_text_pdf()
    create_scanned_pdf()
